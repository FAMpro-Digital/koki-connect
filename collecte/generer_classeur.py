#!/usr/bin/env python3
"""Génère le classeur de collecte des données réelles de Koki (Koki_Collecte_Donnees.xlsx).

Usage : python3 collecte/generer_classeur.py [sortie.xlsx]

Le classeur ne contient que :
  * des numéros d'urgence nationaux, vérifiés auprès d'une source publique citée ;
  * une piste à confirmer (lycée) marquée « Vérifié = Non », donc ignorée à l'import ;
  * la liste des lieux déjà cartographiés dans OpenStreetMap (information).
Aucune autre donnée n'est inventée.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = Path(__file__).resolve().parent.parent
SORTIE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "Koki_Collecte_Donnees.xlsx"

VERT = "0A7D3E"
ORANGE = "E87722"
GRIS = "F2F2F2"
JAUNE = "FFF4CC"

NIVEAUX = ["Préscolaire", "Élémentaire", "Moyen", "Secondaire", "Enseignement religieux",
           "Formation professionnelle", "Supérieur", "Autre"]
SECTEURS = ["Public", "Privé", "Communautaire", "Autre"]
PROFILS = ["Commerçant", "Prestataire de services", "Les deux"]
CATEGORIES_PRO = ["Alimentation", "Maison", "Vêtements", "Électronique", "Agriculture", "Autres produits",
                  "Maçonnerie", "Plomberie", "Électricité", "Menuiserie", "Mécanique", "Couture", "Coiffure",
                  "Soudure", "Peinture", "Carrelage", "Construction", "Réparation téléphone/informatique",
                  "Nettoyage", "Photographie", "Décoration événementielle", "Restauration", "Transport",
                  "Autres métiers"]
GROUPES_NUMEROS = ["Urgences nationales", "Santé", "Sécurité", "Administration", "Eau et électricité",
                   "Transport", "Autre"]
CATEGORIES_LIEUX = ["Santé", "Services publics"]
TYPES_LIEUX = ["Hôpital", "Centre de santé", "Poste de santé", "Pharmacie", "Clinique privée", "Mairie",
               "Sous-préfecture", "Gendarmerie", "Police", "Poste", "Électricité", "Eau", "Justice", "Autre"]
OUI_NON = ["Oui", "Non"]

mince = Side(style="thin", color="BBBBBB")
bordure = Border(left=mince, right=mince, top=mince, bottom=mince)


def feuille(wb, titre, colonnes, lignes=(), validations=None, hauteur_entete=34, nb_lignes_vides=300):
    """colonnes : [(titre, largeur, obligatoire)]. validations : {indice_colonne: liste | ('decimal', min, max) | 'tel' | 'date'}."""
    ws = wb.create_sheet(titre)
    for i, (nom, largeur, obligatoire) in enumerate(colonnes, start=1):
        c = ws.cell(row=1, column=i, value=nom + (" *" if obligatoire else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERT if obligatoire else "5B6B63")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = bordure
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.row_dimensions[1].height = hauteur_entete
    ws.freeze_panes = "A2"
    for r, ligne in enumerate(lignes, start=2):
        for i, v in enumerate(ligne, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = bordure
    derniere = max(len(list(lignes)) + 1, 1) + nb_lignes_vides
    for indice, regle in (validations or {}).items():
        lettre = get_column_letter(indice)
        plage = f"{lettre}2:{lettre}{derniere}"
        if isinstance(regle, list):
            texte = ",".join(regle)
            if len(texte) < 250:
                dv = DataValidation(type="list", formula1=f'"{texte}"', allow_blank=True)
            else:  # liste trop longue : feuille « Listes »
                dv = DataValidation(type="list", formula1=regle_liste(wb, regle), allow_blank=True)
            dv.errorTitle, dv.error = "Valeur non prévue", "Choisissez une valeur dans la liste."
        elif regle == "tel":
            dv = DataValidation(type="custom", allow_blank=True,
                                formula1=f'=AND(LEFT(SUBSTITUTE({lettre}2," ",""),1)="+",LEN(SUBSTITUTE({lettre}2," ",""))>=9,LEN(SUBSTITUTE({lettre}2," ",""))<=16)')
            dv.errorTitle, dv.error = "Téléphone", "Format international attendu : +221 77 123 45 67 (ou un numéro court d'urgence sans +)."
            dv.errorStyle = "warning"
        elif regle == "date":
            dv = DataValidation(type="custom", allow_blank=True,
                                formula1=f'=AND(LEN({lettre}2)=10,MID({lettre}2,5,1)="-",MID({lettre}2,8,1)="-")')
            dv.errorTitle, dv.error = "Date", "Écrivez la date au format AAAA-MM-JJ (ex. 2026-10-10)."
        elif isinstance(regle, tuple) and regle[0] == "decimal":
            dv = DataValidation(type="decimal", operator="between", formula1=str(regle[1]), formula2=str(regle[2]), allow_blank=True)
            dv.errorTitle, dv.error = "Coordonnée", f"Nombre décimal entre {regle[1]} et {regle[2]} (ex. 15.5059)."
        else:
            continue
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(plage)
    ws.sheet_view.zoomScale = 100
    return ws


_listes = {}


def regle_liste(wb, valeurs):
    cle = tuple(valeurs)
    if cle not in _listes:
        ws = wb["Listes"] if "Listes" in wb.sheetnames else wb.create_sheet("Listes")
        col = len(_listes) + 1
        for r, v in enumerate(valeurs, start=1):
            ws.cell(row=r, column=col, value=v)
        _listes[cle] = f"=Listes!${get_column_letter(col)}$1:${get_column_letter(col)}${len(valeurs)}"
    return _listes[cle]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.column_dimensions["A"].width = 120
    texte = [
        ("KOKI CONNECT — Collecte des données réelles de Koki", "titre"),
        ("", None),
        ("À quoi sert ce classeur ?", "h"),
        ("À rassembler les informations locales (numéros utiles, écoles, santé, services publics, commerces et artisans) dans un format que l'application peut lire. Les colonnes marquées d'une étoile (*) sont obligatoires.", None),
        ("", None),
        ("Règles à respecter (important)", "h"),
        ("1. Ne saisissez que des informations destinées à être publiques. Pour un commerce ou un artisan, notez son accord (colonnes « Autorisation de publier » et « Date de l'autorisation »).", None),
        ("2. Pas de numéro privé d'une personne sans son accord clair. Pour les commerces, utilisez le numéro WhatsApp professionnel.", None),
        ("3. Une ligne n'est importée que si « Autorisation de publier = Oui » ET « Vérifié = Oui ». Les autres lignes sont listées dans le rapport mais ignorées.", None),
        ("4. « Vérifié = Oui » signifie qu'une personne a appelé ou s'est rendue sur place, ou qu'une source officielle est citée dans la colonne « Source ».", None),
        ("5. Téléphones au format international : +221 77 123 45 67. Seuls les numéros d'urgence courts (17, 18, 1515…) s'écrivent sans « + ».", None),
        ("6. Coordonnées GPS : latitude et longitude en degrés décimaux avec un point (ex. 15.5059 et -15.9919). Depuis un téléphone : Google Maps ou OpenStreetMap, appui long sur le lieu.", None),
        ("7. Dates au format AAAA-MM-JJ (ex. 2026-10-10).", None),
        ("", None),
        ("Feuilles", "h"),
        ("• Numéros utiles : urgences et services (déjà rempli avec les numéros nationaux d'urgence vérifiés).", None),
        ("• Écoles : établissements scolaires (niveau et secteur selon les listes proposées).", None),
        ("• Lieux et services : santé (poste de santé, pharmacie…) et services publics (mairie, gendarmerie…).", None),
        ("• Commerces et artisans : professionnels d'accord pour figurer dans l'annuaire. Ils seront invités à créer eux-mêmes leur fiche (connexion par SMS), puis validés dans l'admin.", None),
        ("• Carte OSM (info) : lieux déjà présents sur la carte OpenStreetMap ; beaucoup n'ont pas de nom. À compléter directement sur openstreetmap.org (c'est là que l'application les lit).", None),
        ("", None),
        ("Après remplissage", "h"),
        ("Envoyez le classeur à Claude (ou lancez : python3 collecte/convertir_collecte.py Koki_Collecte_Donnees.xlsx). Il produit un rapport de contrôle et les fichiers prêts à importer dans collecte/sortie/.", None),
        ("", None),
        ("Lignes pré-remplies : uniquement les numéros d'urgence nationaux (source citée) et une piste à confirmer (lycée, Vérifié = Non). Rien d'autre n'a été inventé.", None),
    ]
    for r, (t, style) in enumerate(texte, start=1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "titre":
            c.font = Font(bold=True, size=16, color=VERT)
        elif style == "h":
            c.font = Font(bold=True, size=12, color=ORANGE)
    ws.sheet_view.showGridLines = False

    # --- Numéros utiles
    src = "Senego.com, liste « Numéros d'urgence au Sénégal » (consultée le 2026-10-10)"
    numeros = [
        ["Urgences nationales", "Police secours", "17", "Appel gratuit, urgences de police.", src, "2026-10-10", "Oui", "Oui"],
        ["Urgences nationales", "Sapeurs-pompiers", "18", "Incendies, accidents, secours.", src, "2026-10-10", "Oui", "Oui"],
        ["Urgences nationales", "SAMU national", "1515", "Urgences médicales.", src, "2026-10-10", "Oui", "Oui"],
        ["Urgences nationales", "Numéro vert violences basées sur le genre", "800 00 20 20", "Gratuit et confidentiel.", src, "2026-10-10", "Oui", "Oui"],
    ]
    feuille(wb, "Numéros utiles",
            [("Groupe", 22, True), ("Nom du service", 38, True), ("Téléphone", 20, True), ("Description", 48, False),
             ("Source", 50, True), ("Date de vérification (AAAA-MM-JJ)", 18, True), ("Autorisation de publier", 14, True), ("Vérifié", 10, True)],
            numeros, {1: GROUPES_NUMEROS, 3: "tel", 6: "date", 7: OUI_NON, 8: OUI_NON})

    # --- Écoles
    ecoles = [[
        "Lycée d'enseignement général de Koki (nouveau lycée)", "Secondaire", "Public", "Koki", "", "",
        "Nouveau lycée inauguré en octobre 2023 (tournée économique dans la région de Louga). À CONFIRMER : nom exact, adresse, téléphone, effectifs.",
        "primature.sn — Inauguration du nouveau lycée de Koki (22 octobre 2023)", "", "Oui", "Non"]]
    feuille(wb, "Écoles",
            [("Nom de l'établissement", 40, True), ("Niveau", 22, True), ("Secteur", 14, True), ("Adresse / village", 28, True),
             ("Téléphone public", 20, False), ("Horaires", 26, False), ("Description", 48, False), ("Source", 40, True),
             ("Date de vérification (AAAA-MM-JJ)", 18, True), ("Autorisation de publier", 14, True), ("Vérifié", 10, True)],
            ecoles, {2: NIVEAUX, 3: SECTEURS, 5: "tel", 9: "date", 10: OUI_NON, 11: OUI_NON})

    # --- Lieux et services (santé, services publics)
    feuille(wb, "Lieux et services",
            [("Catégorie", 16, True), ("Type", 18, True), ("Nom", 38, True), ("Adresse / village", 28, True),
             ("Téléphone public", 20, False), ("Horaires", 26, False), ("Garde 24 h", 10, False),
             ("Latitude", 12, False), ("Longitude", 12, False), ("Description", 40, False), ("Source", 36, True),
             ("Date de vérification (AAAA-MM-JJ)", 18, True), ("Autorisation de publier", 14, True), ("Vérifié", 10, True)],
            [], {1: CATEGORIES_LIEUX, 2: TYPES_LIEUX, 5: "tel", 7: OUI_NON, 8: ("decimal", -90, 90), 9: ("decimal", -180, 180),
                 12: "date", 13: OUI_NON, 14: OUI_NON})

    # --- Commerces et artisans
    feuille(wb, "Commerces et artisans",
            [("Nom commercial", 34, True), ("Profil", 22, True), ("Catégorie / métier", 26, True),
             ("Téléphone WhatsApp professionnel", 22, True), ("Adresse / localisation", 34, True),
             ("Latitude", 12, True), ("Longitude", 12, True), ("Horaires", 28, False),
             ("Description (10 caractères minimum)", 46, True), ("Responsable (NON publié)", 24, False),
             ("Date de l'autorisation (AAAA-MM-JJ)", 18, True), ("Autorisation de publier", 14, True), ("Vérifié", 10, True)],
            [], {2: PROFILS, 3: CATEGORIES_PRO, 4: "tel", 6: ("decimal", -90, 90), 7: ("decimal", -180, 180),
                 11: "date", 12: OUI_NON, 13: OUI_NON})

    # --- Carte OSM (information)
    carte = json.loads((RACINE / "assets" / "koki-carte.json").read_text(encoding="utf-8"))
    lignes = []
    for lieu in sorted(carte["lieux"], key=lambda x: (x["categorie"], x["nom"])):
        lignes.append([lieu["id"], lieu["categorie"], lieu["nom"], lieu["latitude"], lieu["longitude"],
                       "Oui" if "nom non renseigné" in lieu["nom"] else "Non"])
    ws = feuille(wb, "Carte OSM (info)",
                 [("Identifiant OpenStreetMap", 24, False), ("Catégorie", 22, False), ("Nom actuel", 46, False),
                  ("Latitude", 12, False), ("Longitude", 12, False), ("Nom à compléter ?", 14, False)],
                 lignes, nb_lignes_vides=0)
    for row in ws.iter_rows(min_row=2):
        if row[5].value == "Oui":
            for c in row:
                c.fill = PatternFill("solid", fgColor=JAUNE)
    ws["H1"] = f"Source : OpenStreetMap, relation {carte['relation']}, extraction du {carte['date']}. Ces lieux sont lus directement par la carte de l'application : pour les nommer, modifiez-les sur openstreetmap.org (et non dans ce classeur)."
    ws["H1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["H"].width = 60

    if "Listes" in wb.sheetnames:
        wb["Listes"].sheet_state = "hidden"
    wb.save(SORTIE)
    print(f"Classeur écrit : {SORTIE}")


if __name__ == "__main__":
    main()
